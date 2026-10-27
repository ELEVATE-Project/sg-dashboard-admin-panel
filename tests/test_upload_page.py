from io import BytesIO
from pathlib import Path
import sys
import types
import unittest
from unittest.mock import patch

import openpyxl
from streamlit.testing.v1 import AppTest

from upload_page_module import upload_page

ACTIONS = upload_page.ACTIONS
from processing_report import record_upload, report_print


class UploadPageTests(unittest.TestCase):
    def setUp(self):
        workbook = openpyxl.Workbook()
        workbook.active.title = 'Goals'
        workbook.active.append(['Indicator', 'Data'])
        workbook.active.append(['Schools', 12])
        homepage = workbook.create_sheet('Data on homepage')
        homepage.append(['Indicator', 'Definition', 'Data', 'Icon link'])
        homepage.append(['Schools', 'Schools in the network', 12, ''])
        workbook.create_sheet('Instructions').append(['Not processed'])
        self.upload = BytesIO()
        workbook.save(self.upload)
        self.upload.name = 'multi-sheet.xlsx'
        self.page = Path(__file__).resolve().parents[1] / 'frontend-pages' / 'file-upload.py'
        self.callbacks = {}
        self.archive = {'gcs_path': 'gs://test/source.xlsx'}
        self.modules = {}
        names = [spec.processor for spec in ACTIONS] + [
            'community_led_details.pie_chart_community_led',
            'upload_excel_to_gcs.upload_excel_to_gcs',
        ]
        for name in names:
            module_name, function_name = name.split('.')
            module_name = 'tabs_scripts.' + module_name
            module = self.modules.setdefault(module_name, types.ModuleType(module_name))
            setattr(module, function_name, self.fake_function(function_name))

    def fake_function(self, name):
        def function(source):
            if name == 'upload_excel_to_gcs':
                return self.archive
            if name in self.callbacks:
                return self.callbacks[name](source)
        function.__name__ = name
        return function

    def add_sheets(self, sheets):
        workbook = openpyxl.load_workbook(BytesIO(self.upload.getvalue()))
        for name, rows in sheets.items():
            sheet = workbook.create_sheet(name)
            for row in rows:
                sheet.append(row)
        self.upload = BytesIO()
        workbook.save(self.upload)
        self.upload.name = 'multi-sheet.xlsx'

    def run_sheet(self, sheet, expected_steps):
        calls = []
        for name in expected_steps:
            self.callbacks[name] = lambda source, name=name: calls.append(name)
        with patch('streamlit.file_uploader', return_value=self.upload), patch.dict(sys.modules, self.modules):
            app = AppTest.from_file(str(self.page)).run()
            next(button for button in app.button if button.label == f'Upload {sheet}').click().run()
        self.assertFalse(app.exception)
        self.assertEqual(calls, expected_steps)
        self.assertEqual([item['processor'].split('.')[-1] for item in app.session_state['_excel_results']], expected_steps)
        return app

    def test_homepage_upload_updates_its_dependent_overview(self):
        self.add_sheets({
            'States details': [['State Name', 'Indicator', 'Definition', 'Data'], ['State', 'Schools', '', 1]],
            'State_district details': [['State'], ['State']],
        })
        self.run_sheet('Data on homepage', ['key_progress_indicators', 'update_district_view_indicators'])

    def test_micro_upload_runs_both_chart_steps_in_order(self):
        self.add_sheets({
            'Micro improvements progress': [['State', 'District', 'Year', 'Q1'], ['State', None, 2025, 1]],
            'Graph_VoiceTab_MI': [['State', 'District', 'Year', 'Q1'], ['State', None, 2025, 1]],
        })
        self.run_sheet('Micro improvements progress', ['extract_micro_improvements', 'update_voices_json_line_chart'])

    def test_program_upload_attaches_frameworks_after_regeneration(self):
        self.add_sheets({
            'Programs': [['State Name', 'District Name', 'Program Type', 'Name of the Program'], ['State', '', 'SLC', 'Example']],
            'Impact_Programs': [
                ['State Name', 'District Name', 'Program Type', 'Name of the Program', 'Impact Layer',
                 'Framework Dimension / Change', 'Impact Explanation / Points', 'Evidence Link'],
                ['State', '', 'SLC', 'Example', 'System', 'Change', 'Details', ''],
            ],
        })
        self.run_sheet('Programs', ['generate_program_reports', 'update_impact_program_frameworks'])

    def test_community_upload_runs_totals_and_district_details(self):
        self.add_sheets({
            'New_Community led Programs ': [['State'], ['State']],
            'State_district details': [['State'], ['State']],
        })
        self.run_sheet('New_Community led Programs ', ['community_led_programs_sum_with_codes', 'extract_community_details'])

    def test_missing_followup_sheet_is_reported_and_not_called(self):
        self.add_sheets({'Micro improvements progress': [['State', 'Year'], ['State', 2025]]})
        calls = []
        self.callbacks['extract_micro_improvements'] = lambda source: calls.append('micro')
        self.callbacks['update_voices_json_line_chart'] = lambda source: self.fail('Missing sheet must not be processed')
        with patch('streamlit.file_uploader', return_value=self.upload), patch.dict(sys.modules, self.modules):
            app = AppTest.from_file(str(self.page)).run()
            next(button for button in app.button if button.label == 'Upload Micro improvements progress').click().run()
        self.assertFalse(app.exception)
        self.assertEqual(calls, ['micro'])
        self.assertEqual(app.session_state['_excel_results'][1]['status'], 'Skipped')

    def test_failure_in_first_sheet_step_does_not_hide_followup(self):
        self.add_sheets({
            'States details': [['State Name', 'Indicator', 'Definition', 'Data'], ['State', 'Schools', '', 1]],
            'State_district details': [['State'], ['State']],
        })
        def fail(source):
            raise ValueError('Invalid homepage value')
        self.callbacks['key_progress_indicators'] = fail
        self.callbacks['update_district_view_indicators'] = lambda source: record_upload('overview.json', 'test', 'overview.json')
        with patch('streamlit.file_uploader', return_value=self.upload), patch.dict(sys.modules, self.modules):
            app = AppTest.from_file(str(self.page)).run()
            next(button for button in app.button if button.label == 'Upload Data on homepage').click().run()
        self.assertFalse(app.exception)
        self.assertEqual([item['status'] for item in app.session_state['_excel_results']], ['Failed', 'Uploaded'])

    def test_bulk_report_shows_every_sheet_and_partial_failure_persists(self):
        calls = []
        def homepage(source):
            calls.append('homepage')
            report_print('Error: Invalid data in row 2')
        def goals(source):
            calls.append('goals')
            record_upload('dashboard.json', 'test', 'dashboard.json', payload='{"schools":12}')
        self.callbacks.update(key_progress_indicators=homepage, goals=goals)
        self.callbacks['pie_chart_community_led'] = lambda source: self.fail('Disabled step must stay disabled')
        with patch('streamlit.file_uploader', return_value=self.upload), patch.dict(sys.modules, self.modules):
            app = AppTest.from_file(str(self.page)).run()
            self.assertFalse(app.exception)
            self.assertEqual([tab.label for tab in app.tabs], ['Goals', 'Data on homepage', 'Instructions'])
            self.assertTrue(any('no configured processor' in warning.value for warning in app.warning))
            next(button for button in app.button if 'Upload all files' in button.label).click().run()
            self.assertFalse(app.exception)
            results = app.session_state['_excel_results']
            self.assertEqual([item['status'] for item in results if item['status'] != 'Skipped'], ['Failed', 'Uploaded'])
            self.assertEqual([item['processor'] for item in results], [spec.processor for spec in ACTIONS])
            self.assertTrue(any('need attention' in warning.value for warning in app.warning))
            self.assertFalse(app.success)
            self.assertTrue(app.json)
            app.run()
            self.assertFalse(app.exception)
            self.assertEqual(calls, ['homepage', 'goals'])
            self.assertEqual(len(app.session_state['_excel_results']), len(ACTIONS))
            self.upload.name = 'another-workbook.xlsx'
            app.run()
            self.assertNotIn('_excel_results', app.session_state)
            self.assertNotIn('_excel_archive', app.session_state)

    def test_single_sheet_upload_warns_if_no_updates_are_confirmed(self):
        self.archive = None
        with patch('streamlit.file_uploader', return_value=self.upload), patch.dict(sys.modules, self.modules):
            app = AppTest.from_file(str(self.page)).run()
            next(button for button in app.button if button.label == 'Upload Goals').click().run()
            self.assertFalse(app.exception)
            self.assertFalse(app.success)
            self.assertTrue(any('could not be archived' in warning.value for warning in app.warning))
            self.assertEqual(app.session_state['_excel_results'][0]['status'], 'Warning')
            self.assertEqual(len(app.session_state['_excel_results']), 1)

    def test_retained_file_is_previewed_when_browser_uploader_is_empty(self):
        with patch('streamlit.file_uploader', return_value=self.upload) as uploader, patch.dict(sys.modules, self.modules):
            app = AppTest.from_file(str(self.page)).run()
            uploader.return_value = None
            app.run()
            self.assertFalse(app.exception)
            self.assertTrue(any('Current file: multi-sheet.xlsx' in item.value for item in app.caption))
            self.assertEqual([tab.label for tab in app.tabs], ['Goals', 'Data on homepage', 'Instructions'])
            self.assertNotIn('_excel_results', app.session_state)
            # A newly selected file replaces the retained file.
            self.upload.name = 'latest-workbook.xlsx'
            uploader.return_value = self.upload
            app.run()
            self.assertEqual(app.session_state['_last_uploaded_excel']['name'], 'latest-workbook.xlsx')

    def test_commenting_a_bulk_call_disables_that_step(self):
        code = self.page.read_text().replace(
            '                            goals(uploaded_file)',
            '                            # goals(uploaded_file)',
        )
        self.callbacks['goals'] = lambda source: self.fail('Commented step must not run')
        with patch('streamlit.file_uploader', return_value=self.upload), patch.dict(sys.modules, self.modules):
            app = AppTest.from_string(code).run()
            next(button for button in app.button if 'Upload all files' in button.label).click().run()
        self.assertFalse(app.exception)
        self.assertFalse(any(item['processor'] == 'goals.goals' for item in app.session_state['_excel_results']))

    def test_program_import_error_does_not_prevent_other_sheets(self):
        self.modules['tabs_scripts.programs'] = types.ModuleType('tabs_scripts.programs')
        self.callbacks['goals'] = lambda source: record_upload('dashboard.json', 'test', 'dashboard.json')
        with patch('streamlit.file_uploader', return_value=self.upload), patch.dict(sys.modules, self.modules):
            app = AppTest.from_file(str(self.page)).run()
            next(button for button in app.button if 'Upload all files' in button.label).click().run()
        self.assertFalse(app.exception)
        self.assertTrue(any(item['processor'] == 'goals.goals' and item['status'] == 'Uploaded'
                            for item in app.session_state['_excel_results']))

    def test_csv_has_preview_without_dashboard_upload_action(self):
        self.upload = BytesIO(b'Name,Value\nSchools,12\n')
        self.upload.name = 'data.csv'
        with patch('streamlit.file_uploader', return_value=self.upload):
            app = AppTest.from_file(str(self.page)).run()
        self.assertFalse(app.exception)
        self.assertFalse(app.button)
        self.assertTrue(app.dataframe)
        self.assertIn('requires an Excel', app.info[0].value)

from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
import importlib.util
import json
from pathlib import Path
import tempfile
import types
import unittest
from unittest.mock import MagicMock, patch

import openpyxl

from upload_page_module import upload_page

Action = upload_page.Action
inspect_workbook = upload_page.inspect_workbook
json_changes = upload_page.json_changes
processing_session = upload_page.processing_session
run_action = upload_page.run_action
from processing_report import _current_report, record_upload, report_print


class ExcelProcessingTests(unittest.TestCase):
    def setUp(self):
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.root = Path(directory.name)
        (self.root / 'pages').mkdir()
        self.output = self.root / 'pages' / 'dashboard.json'
        self.output.write_text('{"value": 1}')
        workbook = openpyxl.Workbook()
        workbook.active.title = 'Goals'
        workbook.active.append(['Indicator', 'Data'])
        workbook.active.append(['Schools', 10])
        workbook.create_sheet('Unknown sheet').append(['Unmapped'])
        data = BytesIO()
        workbook.save(data)
        workbook.close()
        self.data = data.getvalue()
        self.inventory = inspect_workbook(self.data)
        self.spec = Action('Goals', 'fake.process', ('Indicator', 'Data'))

    def run_actions(self, actions, data=None, inventory=None):
        data = self.data if data is None else data
        inventory = self.inventory if inventory is None else inventory
        with processing_session():
            return [run_action(spec, data, inventory, self.root) for spec in actions]

    def run_processor(self, function, actions=None):
        with patch.object(upload_page, 'import_module', return_value=types.SimpleNamespace(process=function)):
            return self.run_actions(actions or (self.spec,))

    def test_swallowed_errors_are_failures_and_remaining_steps_run(self):
        calls = []
        def processor(source):
            calls.append(source.tell())
            source.read()
            if len(calls) == 1:
                report_print('Error: Missing required data')
                return
            self.output.write_text('{"value": 10}')
            record_upload(self.output, 'test', 'dashboard.json', payload='{"value": 10}')
        results = self.run_processor(processor, (self.spec, self.spec))
        self.assertEqual([result['status'] for result in results], ['Failed', 'Uploaded'])
        self.assertEqual(calls, [0, 0])
        self.assertEqual(results[1]['outputs'][0]['changes'], [
            {'path': '/value', 'change': 'replaced', 'before': 1, 'after': 10}])
        self.assertEqual(results[1]['uploads'][0]['data'], {'value': 10})
        self.assertIsNone(_current_report.get())

    def test_exceptions_do_not_abort_later_sheets(self):
        attempts = []
        def processor(source):
            attempts.append(1)
            if len(attempts) == 1:
                raise ValueError('Invalid number in row 2')
        results = self.run_processor(processor, (self.spec, self.spec))
        self.assertEqual(len(results), 2)
        self.assertEqual(results[0]['status'], 'Failed')
        self.assertIn('Invalid number in row 2', results[0]['messages'][0]['message'])
        self.assertEqual(results[1]['status'], 'Warning')

    def test_legacy_exit_is_reported_without_aborting_remaining_steps(self):
        calls = []
        def processor(source):
            calls.append(1)
            if len(calls) == 1:
                raise SystemExit(1)
        results = self.run_processor(processor, (self.spec, self.spec))
        self.assertEqual(len(calls), 2)
        self.assertEqual(results[0]['status'], 'Failed')
        self.assertIn('exited with status 1', results[0]['messages'][0]['message'])

    def test_none_return_and_success_text_do_not_prove_publication(self):
        results = self.run_processor(lambda source: report_print('Updated successfully!'))
        self.assertEqual(results[0]['status'], 'Warning')
        self.assertIn('No data changes or cloud uploads', results[0]['summary'])

    def test_local_change_and_cloud_failure_are_partial(self):
        def processor(source):
            self.output.write_text('{"value": 10}')
            record_upload(self.output, 'test', 'dashboard.json', error='Permission denied')
        result = self.run_processor(processor)[0]
        self.assertEqual(result['status'], 'Partial')
        self.assertEqual(result['uploads'][0]['status'], 'failed')
        self.assertEqual(result['outputs'][0]['data'], {'value': 10})

    def test_local_only_changes_do_not_count_as_uploaded(self):
        def processor(source):
            self.output.write_text('{"value": 2}')
        result = self.run_processor(processor)[0]
        self.assertEqual(result['status'], 'Warning')
        self.assertIn('no cloud upload was confirmed', result['summary'])

    def test_unchanged_json_can_be_confirmed_uploaded(self):
        result = self.run_processor(lambda source: record_upload(
            self.output, 'test', 'dashboard.json', payload='{"value": 1}'))[0]
        self.assertEqual(result['status'], 'Uploaded')
        self.assertFalse(result['outputs'])
        self.assertIn('0 local JSON file(s) changed', result['summary'])

    def test_preflight_missing_empty_wrong_name_columns_and_dependencies(self):
        inventory = self.inventory
        inventory['Empty'] = {'rows': 0, 'headers': []}
        inventory['Wrong columns'] = {'rows': 1, 'headers': ['Other']}
        inventory[' goals '] = inventory.pop('Goals')
        specs = (
            Action('Missing', 'fake.process'), Action('Empty', 'fake.process'),
            self.spec, Action('Wrong columns', 'fake.process', ('Required',)),
            Action(' goals ', 'fake.process', dependencies=('Supporting sheet',)),
        )
        with patch.object(upload_page, 'import_module') as loader:
            results = self.run_actions(specs, inventory=inventory)
        loader.assert_not_called()
        self.assertEqual([item['status'] for item in results], ['Skipped', 'Skipped', 'Failed', 'Failed', 'Failed'])
        self.assertIn('Rename sheet', results[2]['summary'])
        self.assertIn('Required', results[3]['summary'])
        self.assertIn('Supporting sheet', results[4]['summary'])

    def test_inventory_counts_rows_and_includes_unmapped_sheets(self):
        self.assertEqual(self.inventory['Goals']['rows'], 1)
        self.assertIn('Unknown sheet', self.inventory)

    def test_empty_required_support_sheet_does_not_clear_existing_data(self):
        self.inventory['Supporting'] = {'rows': 0, 'headers': []}
        with patch.object(upload_page, 'import_module') as loader:
            result = self.run_actions((Action('Goals', 'fake.process', dependencies=('Supporting',)),))[0]
        loader.assert_not_called()
        self.assertEqual(result['status'], 'Failed')
        self.assertIn('Supporting sheets have no data rows', result['summary'])
        self.assertEqual(self.output.read_text(), '{"value": 1}')

    def test_context_collection_is_isolated_across_threads(self):
        def collect(message):
            report = {'messages': [], 'uploads': []}
            token = _current_report.set(report)
            try:
                report_print(message)
                return report['messages']
            finally:
                _current_report.reset(token)
        with ThreadPoolExecutor(max_workers=2) as pool:
            first, second = list(pool.map(collect, ['first', 'second']))
        self.assertEqual(first[0]['message'], 'first')
        self.assertEqual(second[0]['message'], 'second')

    def test_diffs_include_additions_removals_and_type_changes(self):
        changes = list(json_changes({'a': True, 'removed': 1}, {'a': 1, 'new': 2}))
        self.assertEqual({item['change'] for item in changes}, {'added', 'removed', 'replaced'})

    def test_errors_after_the_diagnostic_limit_still_fail_the_step(self):
        def processor(source):
            for _ in range(1000):
                report_print('Processing row')
            report_print('Error: final validation failed')
        with patch('processing_report.builtins.print'):
            result = self.run_processor(processor)[0]
        self.assertEqual(result['status'], 'Failed')
        self.assertEqual(result['omitted_messages'], 1)

    def test_actual_goals_missing_columns_is_reported_without_cloud_calls(self):
        # Bypass preflight to exercise the legacy processor's caught error path.
        from tabs_scripts.goals import goals
        bad_workbook = openpyxl.Workbook()
        bad_workbook.active.title = 'Goals'
        bad_workbook.active.append(['Incorrect'])
        bad_workbook.active.append(['value'])
        stream = BytesIO()
        bad_workbook.save(stream)
        data = stream.getvalue()
        with patch.object(upload_page, 'import_module', return_value=types.SimpleNamespace(process=goals)):
            result = self.run_actions((Action('Goals', 'fake.process'),), data=data, inventory=inspect_workbook(data))[0]
        self.assertEqual(result['status'], 'Failed')
        self.assertTrue(any('must contain columns' in message['message'] for message in result['messages']))

    def test_cloud_receipts_require_successful_upload_and_publication(self):
        module_path = Path(__file__).resolve().parents[1] / 'cloud-scripts' / 'gcp_access.py'
        spec = importlib.util.spec_from_file_location('test_gcp_access', module_path)
        gcp = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(gcp)
        for error in (None, RuntimeError('Storage unavailable')):
            with self.subTest(error=error):
                blob = MagicMock()
                blob.public_url = 'https://storage.googleapis.com/test/dashboard.json'
                blob.upload_from_string.side_effect = error
                client = MagicMock()
                client.bucket.return_value.blob.return_value = blob
                report = {'messages': [], 'uploads': []}
                token = _current_report.set(report)
                try:
                    with patch.object(gcp, 'get_storage_client', return_value=client):
                        result = gcp.upload_file_to_gcs_and_get_directory('test', self.output, 'dashboard.json')
                finally:
                    _current_report.reset(token)
                self.assertEqual(report['uploads'][0]['status'], 'failed' if error else 'uploaded')
                self.assertEqual(bool(result), error is None)

import importlib.util
import json
import os
import re

import openpyxl
from dotenv import load_dotenv
from google.oauth2 import service_account
from googleapiclient.discovery import build

from constants import PAGE_METADATA, TABS_METADATA


load_dotenv()

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_ROOT = os.path.dirname(SCRIPT_DIR)
PROGRAM_JSON_FILENAMES = {"state-program.json", "SLC.json", "WLC.json"}
STATE_CODE_PATH = os.path.join(PROJECT_ROOT, "pages", "state_code_details.json")
DRIVE_READONLY_SCOPE = "https://www.googleapis.com/auth/drive.readonly"
drive_service = None
failed_evidence_file_ids = set()


def normalize(text):
    return re.sub(r"[^a-z0-9]", "", str(text or "").strip().lower())


def snake_case(text):
    text = re.sub(r"\s+", "_", str(text or "").strip())
    text = re.sub(r"[^a-zA-Z0-9_]", "", text)
    return text.lower()


def clean_text(value):
    lines = []
    for line in str(value or "").splitlines():
        line = re.sub(r"[ \t]+", " ", line).strip()
        if line:
            lines.append(line)
    return "\n".join(lines)


def clean_json_value(value):
    if isinstance(value, str):
        return clean_text(value)
    return value


def slug_key(text):
    text = re.sub(r"[^a-z0-9]+", "_", str(text or "").strip().lower())
    return text.strip("_")


def clean_url(url):
    url = str(url or "").strip()
    markdown_link_match = re.match(r"^\[([^\]]+)\]\(([^)]+)\)$", url)
    if markdown_link_match:
        return markdown_link_match.group(2).strip()
    return url


def extract_drive_file_id(url):
    url = clean_url(url)
    patterns = [
        r"/file/d/([a-zA-Z0-9_-]+)",
        r"[?&]id=([a-zA-Z0-9_-]+)",
        r"/d/([a-zA-Z0-9_-]+)",
    ]
    for pattern in patterns:
        match = re.search(pattern, url)
        if match:
            return match.group(1)
    return ""


def format_file_size(size_bytes):
    if size_bytes in (None, ""):
        return ""

    size = float(size_bytes)
    units = ["B", "KB", "MB", "GB", "TB"]
    for unit in units:
        if size < 1024 or unit == units[-1]:
            return f"{size:.1f} {unit}" if unit != "B" else f"{int(size)} {unit}"
        size /= 1024

    return ""


def get_file_type(metadata):
    mime_type = metadata.get("mimeType", "")
    file_name = metadata.get("name", "")

    if mime_type == "application/pdf":
        return "PDF"
    if mime_type.startswith("image/"):
        return mime_type.split("/", 1)[1].upper()
    if mime_type.startswith("video/"):
        return mime_type.split("/", 1)[1].upper()
    if mime_type.startswith("audio/"):
        return mime_type.split("/", 1)[1].upper()

    _, extension = os.path.splitext(file_name)
    if extension:
        return extension.lstrip(".").upper()

    if mime_type:
        return mime_type.split("/")[-1].upper()

    return ""


def get_drive_service():
    global drive_service
    if drive_service is not None:
        return drive_service

    private_key = os.getenv("PRIVATE_KEY")
    if not private_key:
        return None

    service_account_info = {
        "type": os.getenv("TYPE"),
        "project_id": os.getenv("PROJECT_ID"),
        "private_key_id": os.getenv("PRIVATE_KEY_ID"),
        "private_key": private_key.replace("\\n", "\n"),
        "client_email": os.getenv("CLIENT_EMAIL"),
        "auth_uri": os.getenv("AUTH_URI"),
        "token_uri": os.getenv("TOKEN_URI"),
        "auth_provider_x509_cert_url": os.getenv("AUTH_PROVIDER_X509_CERT_URL"),
        "client_x509_cert_url": os.getenv("CLIENT_X509_CERT_URL"),
        "universe_domain": os.getenv("UNIVERSE_DOMAIN"),
    }

    try:
        credentials = service_account.Credentials.from_service_account_info(
            service_account_info,
            scopes=[DRIVE_READONLY_SCOPE],
        )
        drive_service = build("drive", "v3", credentials=credentials)
        return drive_service
    except Exception as e:
        print(f"⚠️ Could not initialize Drive metadata lookup: {e}")
        return None


def get_evidence_metadata(evidence_link):
    file_id = extract_drive_file_id(evidence_link)
    if not file_id:
        return {}

    if file_id in failed_evidence_file_ids:
        return None

    service = get_drive_service()
    if not service:
        return None

    try:
        metadata = service.files().get(
            fileId=file_id,
            fields="name,size,mimeType",
            supportsAllDrives=True,
        ).execute()
    except Exception as e:
        failed_evidence_file_ids.add(file_id)
        print(f"⚠️ Could not fetch evidence metadata for {file_id}: {e}")
        return None

    evidence_metadata = {}
    if metadata.get("name"):
        evidence_metadata["evidence_name"] = metadata["name"]
    if metadata.get("size"):
        evidence_metadata["evidence_size_bytes"] = int(metadata["size"])
        evidence_metadata["evidence_size"] = format_file_size(metadata["size"])
    if metadata.get("mimeType"):
        evidence_metadata["evidence_mime_type"] = metadata["mimeType"]
    file_type = get_file_type(metadata)
    if file_type:
        evidence_metadata["evidence_type"] = file_type
    return evidence_metadata


LAYER_KEYS = {
    "learnersoutcomes": "students",
    "schoolsanganwadicentres": "schools",
    "schoolsandanganwadicentres": "schools",
    "community": "community",
    "systeminstitutions": "system",
}


def get_cell_value(cell):
    if cell.hyperlink and cell.hyperlink.target:
        return cell.hyperlink.target
    return cell.value


def get_row_dict(row, header_index_map, expected_columns):
    row_dict = {}
    for column in expected_columns:
        column_index = header_index_map.get(column)
        row_dict[snake_case(column)] = (
            clean_json_value(get_cell_value(row[column_index])) if column_index is not None else ""
        )
    return row_dict


def get_program_match_key(state, district, program_type, program_name):
    return (
        normalize(state),
        normalize(district),
        normalize(program_type),
        normalize(program_name),
    )


def get_program_name_key(state, program_name):
    return (
        normalize(state),
        normalize(program_name),
    )


def load_state_code_map():
    with open(STATE_CODE_PATH, "r", encoding="utf-8") as state_code_file:
        return json.load(state_code_file)


def resolve_state_code(normalized_state, state_code_map):
    for state_name, state_info in state_code_map.items():
        if normalize(state_name) == normalized_state:
            return str(state_info.get("id", "")).strip()
    return ""


def resolve_district_code(normalized_state, normalized_district, state_code_map):
    if not normalized_district:
        return ""

    for state_name, state_info in state_code_map.items():
        if normalize(state_name) != normalized_state:
            continue

        for district_name, district_code in state_info.items():
            if district_name == "id":
                continue
            if normalize(district_name) == normalized_district:
                return str(district_code).strip()

    return ""


def get_framework_row_key(row_dict):
    state = row_dict.get("state_name")
    district = row_dict.get("district_name")
    program_type = row_dict.get("program_type")

    if normalize(program_type) == "state" or normalize(district) == normalize(state):
        district = ""

    return get_program_match_key(
        state,
        district,
        program_type,
        row_dict.get("name_of_the_program"),
    )


def add_framework_entry(framework_by_layer, row_dict):
    impact_layer = str(row_dict.get("impact_layer") or "").strip()
    dimension = str(row_dict.get("framework_dimension__change") or "").strip()
    explanation = str(row_dict.get("impact_explanation__points") or "").strip()
    evidence_link = str(row_dict.get("evidence_link") or "").strip()

    if not (impact_layer or dimension or explanation or evidence_link):
        return

    layer = framework_by_layer.setdefault(
        impact_layer,
        {
            "impact_layer": impact_layer,
            "frameworks": {},
        },
    )
    framework = layer["frameworks"].setdefault(
        dimension,
        {
            "framework_name": dimension,
            "details": [],
        },
    )

    detail = {}
    if explanation:
        detail["description"] = explanation
    if evidence_link:
        detail["evidence_link"] = evidence_link

    if detail:
        framework["details"].append(detail)


def serialize_framework(framework_by_layer):
    return [
        build_framework_layer_payload({
            "impact_layer": layer["impact_layer"],
            "frameworks": list(layer["frameworks"].values()),
        })
        for layer in framework_by_layer.values()
    ]


def get_layer_key(impact_layer):
    normalized_layer = normalize(impact_layer)
    return LAYER_KEYS.get(normalized_layer, slug_key(impact_layer))


def build_framework_layer_payload(layer):
    cards = []
    evidences = []

    for framework in layer.get("frameworks", []):
        framework_name = framework.get("framework_name", "")
        framework_key = slug_key(framework_name)

        for detail in framework.get("details", []):
            description = str(detail.get("description") or "").strip()
            evidence_link = clean_url(detail.get("evidence_link"))

            if description:
                cards.append({
                    "key": framework_key,
                    "label": framework_name,
                    "description": description,
                })

            if evidence_link:
                evidence_metadata = get_evidence_metadata(evidence_link)
                if evidence_metadata is None:
                    continue

                evidence = {
                    "key": framework_key,
                    "title": framework_name,
                    "tag": framework_name,
                    "url": evidence_link,
                }
                evidence.update(evidence_metadata)
                evidences.append(evidence)

    return {
        "layerKey": get_layer_key(layer.get("impact_layer")),
        "impact_layer": layer.get("impact_layer", ""),
        "cards": cards,
        "evidences": evidences,
    }


def build_impact_framework_index(workbook):
    sheet_name = PAGE_METADATA["IMPACT_PROGRAMS"]
    if sheet_name not in workbook.sheetnames:
        print(f"⚠️ Sheet not found: {sheet_name}. Skipping impact program framework.")
        return {}

    sheet = workbook[sheet_name]
    headers = [str(cell.value).strip() if cell.value else "" for cell in sheet[1]]
    header_index_map = {header: index for index, header in enumerate(headers) if header}
    expected_columns = TABS_METADATA["IMPACT_PROGRAMS"]

    missing_columns = [column for column in expected_columns if column not in header_index_map]
    if missing_columns:
        print(f"⚠️ Missing Impact_Programs columns: {missing_columns}. Skipping framework.")
        return {}

    framework_index = {}
    previous_key_values = {
        "state_name": "",
        "district_name": "",
        "program_type": "",
        "name_of_the_program": "",
    }

    for row in sheet.iter_rows(min_row=2):
        row_dict = get_row_dict(row, header_index_map, expected_columns)

        for key in previous_key_values:
            if row_dict.get(key) not in (None, ""):
                previous_key_values[key] = row_dict[key]
            else:
                row_dict[key] = previous_key_values[key]

        program_key = get_framework_row_key(row_dict)
        if not any(program_key):
            continue

        framework_by_layer = framework_index.setdefault(program_key, {})
        add_framework_entry(framework_by_layer, row_dict)

    return {
        program_key: serialize_framework(framework_by_layer)
        for program_key, framework_by_layer in framework_index.items()
    }


def get_program_framework(row_dict, framework_index):
    state = row_dict.get("state_name")
    district = row_dict.get("district_name")
    program_type = row_dict.get("program_type")
    program_name = row_dict.get("name_of_the_program")
    keys = [
        get_program_match_key(
            state,
            district,
            program_type,
            program_name,
        ),
        get_program_match_key(
            state,
            "",
            program_type,
            program_name,
        ),
        get_program_match_key(
            state,
            state,
            program_type,
            program_name,
        ),
    ]

    for program_key in keys:
        framework = framework_index.get(program_key)
        if framework:
            return framework

    program_name_key = get_program_name_key(state, program_name)
    fallback_matches = []
    for indexed_key, framework in framework_index.items():
        indexed_state, indexed_district, indexed_program_type, indexed_program_name = indexed_key
        if (indexed_state, indexed_program_name) != program_name_key:
            continue

        score = 0
        if normalize(district) and indexed_district == normalize(district):
            score += 1
        if normalize(program_type) and indexed_program_type == normalize(program_type):
            score += 1
        fallback_matches.append((score, framework))

    if fallback_matches:
        fallback_matches.sort(key=lambda item: item[0], reverse=True)
        return fallback_matches[0][1]

    return []


def attach_program_framework(row_dict, framework_index):
    row_dict["framework"] = get_program_framework(row_dict, framework_index)
    return row_dict


def get_program_key_from_row(row_dict):
    return get_program_match_key(
        row_dict.get("state_name"),
        row_dict.get("district_name"),
        row_dict.get("program_type"),
        row_dict.get("name_of_the_program"),
    )


def iter_program_json_paths():
    for base_dir in ("states", "districts"):
        root_dir = os.path.join(PROJECT_ROOT, base_dir)
        if not os.path.isdir(root_dir):
            continue

        for current_root, _, filenames in os.walk(root_dir):
            for filename in filenames:
                if filename in PROGRAM_JSON_FILENAMES:
                    yield os.path.join(current_root, filename)


def get_destination_blob_name(json_path):
    relative_path = os.path.relpath(json_path, PROJECT_ROOT)
    return f"sg-dashboard/{relative_path}"


def load_gcp_access():
    gcp_access_path = os.path.join(PROJECT_ROOT, "cloud-scripts", "gcp_access.py")
    spec = importlib.util.spec_from_file_location("gcp_access", gcp_access_path)
    gcp_access = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(gcp_access)
    return gcp_access


def upload_program_json(json_path, gcp_access, bucket_name):
    if not bucket_name:
        print(f"⚠️ BUCKET_NAME not set. Skipping upload for {json_path}.")
        return None

    return gcp_access.upload_file_to_gcs_and_get_directory(
        bucket_name=bucket_name,
        source_file_path=json_path,
        destination_blob_name=get_destination_blob_name(json_path),
    )


def get_candidate_program_json_paths(framework_key, state_code_map):
    normalized_state, normalized_district, normalized_program_type, _ = framework_key
    state_code = resolve_state_code(normalized_state, state_code_map)
    district_code = resolve_district_code(normalized_state, normalized_district, state_code_map)
    paths = []

    if normalized_program_type == "state" and state_code:
        paths.append(os.path.join(PROJECT_ROOT, "states", state_code, "state-program.json"))
    elif normalized_program_type == "slc" and district_code:
        paths.append(os.path.join(PROJECT_ROOT, "districts", district_code, "SLC.json"))
    elif normalized_program_type == "wlc":
        if district_code:
            paths.append(os.path.join(PROJECT_ROOT, "districts", district_code, "WLC.json"))
        if state_code:
            paths.append(os.path.join(PROJECT_ROOT, "states", state_code, "WLC.json"))
    else:
        if district_code:
            paths.append(os.path.join(PROJECT_ROOT, "districts", district_code, f"{normalized_program_type.upper()}.json"))
        if state_code:
            paths.append(os.path.join(PROJECT_ROOT, "states", state_code, "state-program.json"))

    return paths


def update_framework_in_json_file(json_path, framework_key, framework):
    with open(json_path, "r", encoding="utf-8") as json_file:
        programs = json.load(json_file)

    if not isinstance(programs, list):
        return False, 0

    changed = False
    matched_programs = 0
    expected_state, _, expected_program_type, expected_program_name = framework_key

    for program in programs:
        if not isinstance(program, dict):
            continue

        if normalize(program.get("state_name")) != expected_state:
            continue
        if normalize(program.get("name_of_the_program")) != expected_program_name:
            continue
        if expected_program_type and normalize(program.get("program_type")) != expected_program_type:
            continue

        matched_programs += 1
        if program.get("framework") != framework:
            program["framework"] = framework
            changed = True

    if changed:
        with open(json_path, "w", encoding="utf-8") as json_file:
            json.dump(programs, json_file, indent=2, ensure_ascii=False)

    return changed, matched_programs


def get_program_names_from_json(json_path):
    with open(json_path, "r", encoding="utf-8") as json_file:
        programs = json.load(json_file)

    if not isinstance(programs, list):
        return []

    return [
        str(program.get("name_of_the_program") or "").strip()
        for program in programs
        if isinstance(program, dict) and program.get("name_of_the_program")
    ]


def update_program_json_frameworks(framework_index, upload_to_gcs=True):
    updated_files = []
    matched_programs = 0
    matched_framework_keys = set()
    state_code_map = load_state_code_map()

    gcp_access = None
    bucket_name = os.environ.get("BUCKET_NAME")
    if upload_to_gcs:
        try:
            gcp_access = load_gcp_access()
        except Exception as e:
            print(f"⚠️ Could not load GCS helper. Continuing local update only: {e}")

    for framework_key, framework in framework_index.items():
        candidate_paths = get_candidate_program_json_paths(framework_key, state_code_map)
        existing_candidate_paths = [path for path in candidate_paths if os.path.exists(path)]

        if not existing_candidate_paths:
            print(
                "⚠️ No target program JSON found for "
                f"state={framework_key[0] or '-'}, district={framework_key[1] or '-'}, "
                f"type={framework_key[2] or '-'}, program={framework_key[3] or '-'}"
            )
            continue

        key_matched = False
        for json_path in existing_candidate_paths:
            changed, file_matches = update_framework_in_json_file(json_path, framework_key, framework)
            if not file_matches:
                continue

            key_matched = True
            matched_programs += file_matches
            matched_framework_keys.add((framework_key[0], framework_key[3]))

            if not changed:
                continue

            if json_path not in updated_files:
                updated_files.append(json_path)
            print(f"✅ Updated framework in {json_path}")

        if not key_matched:
            available_programs = []
            for json_path in existing_candidate_paths:
                available_programs.extend(get_program_names_from_json(json_path))
            print(
                "⚠️ Target JSON found but program not matched inside file for "
                f"state={framework_key[0] or '-'}, district={framework_key[1] or '-'}, "
                f"type={framework_key[2] or '-'}, program={framework_key[3] or '-'}"
            )
            if available_programs:
                print(f"  Available programs: {available_programs[:10]}")

    for json_path in updated_files:
        if not gcp_access:
            continue

        try:
            folder_url = upload_program_json(json_path, gcp_access, bucket_name)
            if folder_url:
                print(f"✅ Uploaded updated program JSON to {folder_url}")
            else:
                print(f"❌ Failed to upload updated program JSON: {json_path}")
        except Exception as e:
            print(f"❌ Error uploading updated program JSON {json_path}: {e}")

    unmatched_frameworks = []
    for indexed_key in framework_index:
        indexed_state, _, _, indexed_program_name = indexed_key
        if (indexed_state, indexed_program_name) not in matched_framework_keys:
            unmatched_frameworks.append(indexed_key)

    if unmatched_frameworks:
        print("⚠️ Impact framework rows without matching program JSON:")
        for state, district, program_type, program_name in unmatched_frameworks[:20]:
            print(
                "  - "
                f"state={state or '-'}, district={district or '-'}, "
                f"type={program_type or '-'}, program={program_name or '-'}"
            )
        if len(unmatched_frameworks) > 20:
            print(f"  ... and {len(unmatched_frameworks) - 20} more")

    return {
        "updated_files": updated_files,
        "matched_programs": matched_programs,
        "framework_programs": len(framework_index),
        "unmatched_frameworks": unmatched_frameworks,
    }


def update_impact_program_frameworks(excel_file, upload_to_gcs=True):
    workbook = openpyxl.load_workbook(excel_file, data_only=True)
    framework_index = build_impact_framework_index(workbook)
    if not framework_index:
        return {
            "updated_files": [],
            "matched_programs": 0,
            "framework_programs": 0,
            "unmatched_frameworks": [],
        }

    result = update_program_json_frameworks(framework_index, upload_to_gcs=upload_to_gcs)
    print(
        "✅ Impact program framework update complete: "
        f"{len(result['updated_files'])} files updated, "
        f"{result['matched_programs']} programs matched."
    )
    return result

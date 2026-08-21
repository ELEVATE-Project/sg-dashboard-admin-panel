import argparse
import copy
import importlib.util
import json
import os
import re

import openpyxl
from dotenv import load_dotenv

from constants import GCS_STORAGE_BASE_URL


load_dotenv()


SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_ROOT = os.path.dirname(SCRIPT_DIR)
OUTPUT_FILE = os.path.join(PROJECT_ROOT, "pages", "outcomes-model-config.json")
DEFAULT_SHEET_NAME = "Content Requirements"
CONTENT_START_ROW = 60

LAYER_ACTIONS = {
    "learner": "students",
    "schoolandaanganwadi": "schools",
    "schoolandanganwadi": "schools",
    "community": "community",
    "society": "society",
    "systeminstitution": "system",
    "systeminstitutions": "system",
    "network": "network",
}

DEFAULT_LIST_ICONS = {
    "students": ["fact_check", "accessibility_new", "menu_book", "star", "favorite_border"],
    "schools": ["workspace_premium", "menu_book", "co_present"],
    "community": ["groups", "workspace_premium"],
    "system": ["account_balance_wallet", "rate_review", "emoji_objects", "hub"],
    "society": [],
    "network": [],
}

OUTCOMES_MODEL_TEMPLATE = {
    "layerHeading": "Impact Narrative",
    "title": "What Shapes Student Outcomes",
    "description": "A learner's education is shaped across multiple layers:",
    "layerFootnote": "Click on any of the above Layer to read more about it.",
    "chipFootnote": "Click on any of the above quick chip to read more about it.",
    "defaultLayer": "students",
    "programChipLabel": "Learner Outcomes:",
    "programColor": "#ff9911",
    "frameworkHeading": "Impact Framework",
    "frameworkTitle": "Measuring Impact Across Layers",
    "frameworkLead": (
        "The Shikshagraha movement measures impact across these interconnected layers, "
        "each a level through which micro-improvements contribute to systemic transformation."
    ),
    "evidenceHeading": "Evidences and Resources",
    "programChipFootnote": "* Click on any of the above tile to read more about it.",
    "defaultNarrativeBody": "Tap a layer to see what it means.",
    "defaultCtaLabel": "Know more",
    "defaultPartnerImage": "assets/partners/default-partner.svg",
    "defaultCardLabelPrefix": "Card ",
    "ctaIcon": "north_east",
    "staticTexts": {
        "programModeHeading": "Program Details",
        "programModeDescription": "View program outcomes and evidence",
        "diagramNotePrefix": "* ",
        "chipNotePrefix": "* ",
        "evidenceCountLabel": "Evidences and Resources",
        "infoModalTitlePrefix": "About",
        "infoModalDescriptionPrefix": "Information about",
    },
    "ariaLabels": {
        "diagram": "Outcome layers",
        "learner": "Learner",
        "school": "School and Anganwadi",
        "community": "Community",
        "society": "Society",
        "system": "System Institutions",
        "network": "Network",
        "layerTiles": "Outcome layer tiles",
        "narrativeFilters": "Outcome narrative layer filters",
        "frameworkFilters": "Outcome framework layer filters",
        "cardPages": "Card pages",
        "prevCards": "Show previous cards",
        "nextCards": "Show next cards",
        "prevEvidence": "Show previous evidences",
        "nextEvidence": "Show next evidences",
        "viewEvidence": "View evidence",
        "closeModal": "Close",
    },
    "layers": [
        {
            "key": "students",
            "chipLabel": "Learner",
            "diagramLabel": "Learner",
            "icon": "child_care",
            "color": "#ff9911",
            "fill": "#fff3e2",
            "diagram": {
                "shape": "full",
                "innerRadius": 0,
                "outerRadius": 58,
                "labelX": 300,
                "labelY": 300,
                "labelLayout": "icon-only",
                "labelAnchor": "middle",
                "iconX": 275,
                "iconY": 275,
                "dataLayerAttr": "learner",
                "icon": {
                    "type": "material",
                    "value": "assets/icons/learner.svg",
                    "color": "#ff9911",
                    "width": 50,
                    "height": 34,
                },
                "text": {
                    "x": 300,
                    "y": 325,
                    "fill": "#1a1622",
                    "fontSize": 23,
                    "fontWeight": 600,
                },
            },
            "panelType": "list",
            "imgPath": "assets/icons/student.svg",
            "heading": "Learners Outcomes",
            "subheading": "",
            "body": (
                "The learner at the centre - their learning, well-being and aspirations "
                "are what every layer works toward."
            ),
            "listItems": [],
        },
        {
            "key": "schools",
            "chipLabel": "School & Anganwadi Centres",
            "diagramLabel": "School & Anganwadi",
            "icon": "account_balance",
            "color": "#5b6ee0",
            "fill": "#eff0fc",
            "diagram": {
                "shape": "top",
                "innerRadius": 58,
                "outerRadius": 152,
                "labelX": 300,
                "labelY": 402,
                "labelLayout": "stacked",
                "labelAnchor": "middle",
                "iconOffsetY": -18,
                "textOffsetY": 20,
                "iconX": 277,
                "iconY": 366,
                "curvedLabelPath": "M167.43,341.80 A139,139 0 0 0 432.57,341.80",
                "labelForAttr": "school & anganwadi",
                "icon": {
                    "type": "material",
                    "value": "assets/icons/school.svg",
                    "color": "#5b6ee0",
                    "width": 50,
                    "height": 34,
                },
                "text": {
                    "x": 300,
                    "y": 402,
                    "fill": "#1a1622",
                    "fontSize": 23,
                    "fontWeight": 600,
                },
            },
            "panelType": "list",
            "imgPath": "assets/icons/school-solid.svg",
            "heading": "Schools and Anganwadi Improvement",
            "subheading": "",
            "body": (
                "Schools and anganwadi centres - where teaching happens and children spend "
                "most of their day."
            ),
            "listItems": [],
        },
        {
            "key": "community",
            "chipLabel": "Community",
            "diagramLabel": "Community",
            "icon": "group",
            "color": "#e0338a",
            "fill": "#fcebf3",
            "diagram": {
                "shape": "bottom",
                "innerRadius": 58,
                "outerRadius": 152,
                "labelX": 300,
                "labelY": 198,
                "labelLayout": "stacked",
                "labelAnchor": "middle",
                "iconOffsetY": -18,
                "textOffsetY": 22,
                "iconX": 277,
                "iconY": 188,
                "curvedLabelPath": "M219.82,209.38 A121,121 0 0 1 380.18,209.38",
                "icon": {
                    "type": "material",
                    "value": "assets/icons/community.svg",
                    "color": "#e0338a",
                    "width": 50,
                    "height": 34,
                },
                "text": {
                    "x": 300,
                    "y": 198,
                    "fill": "#1a1622",
                    "fontSize": 23,
                    "fontWeight": 600,
                },
            },
            "panelType": "story",
            "imgPath": "assets/icons/community-group.svg",
            "eyebrow": "Community",
            "heading": "Community",
            "subheading": "",
            "body": (
                "Families and local networks that support and demand good education for "
                "their children."
            ),
            "listItems": [],
        },
        {
            "key": "society",
            "chipLabel": "Society",
            "diagramLabel": "Society",
            "icon": "diversity_3",
            "color": "#99459a",
            "fill": "#f2e7f2",
            "diagram": {
                "shape": "top",
                "innerRadius": 152,
                "outerRadius": 216,
                "labelX": 300,
                "labelY": 118,
                "labelLayout": "inline",
                "labelAnchor": "start",
                "textOffsetX": 34,
                "textOffsetY": 4,
                "iconX": 235,
                "iconY": 100,
                "curvedLabelPath": "M233.70,127.29 A185,185 0 0 1 366.30,127.29",
                "icon": {
                    "type": "material",
                    "value": "assets/icons/society.svg",
                    "color": "#99459a",
                    "width": 34,
                    "height": 34,
                },
                "text": {
                    "x": 300,
                    "y": 118,
                    "fill": "#1a1622",
                    "fontSize": 23,
                    "fontWeight": 600,
                },
            },
            "panelType": "story",
            "imgPath": "",
            "eyebrow": "Society",
            "body": (
                "The broader social norms and structures that surround the community."
            ),
            "frameworkNote": "Framework being defined.",
            "cta": {
                "label": "Explore Voices on the Ground",
                "link": "/voices-from-the-ground",
            },
            "frameworkCta": {
                "label": "Explore Voices from the Ground",
                "link": "/voices-from-the-ground",
            },
        },
        {
            "key": "system",
            "chipLabel": "System Institutions",
            "diagramLabel": "System Institutions",
            "icon": "article",
            "color": "#562f91",
            "fill": "#e9e4f1",
            "diagram": {
                "shape": "bottom",
                "innerRadius": 152,
                "outerRadius": 216,
                "labelX": 300,
                "labelY": 482,
                "labelLayout": "inline",
                "labelAnchor": "start",
                "textOffsetX": 34,
                "textOffsetY": 4,
                "iconX": 195,
                "iconY": 460,
                "curvedLabelPath": "M138.95,423.58 A203,203 0 0 0 461.05,423.58",
                "labelForAttr": "system institutions",
                "icon": {
                    "type": "material",
                    "value": "assets/icons/system-institutions.svg",
                    "color": "#562f91",
                    "width": 34,
                    "height": 34,
                },
                "text": {
                    "x": 300,
                    "y": 482,
                    "fill": "#1a1622",
                    "fontSize": 23,
                    "fontWeight": 600,
                },
            },
            "panelType": "list",
            "imgPath": "",
            "eyebrow": "System Institutions",
            "heading": "System Institutions",
            "subheading": "Strengthening public education systems",
            "body": (
                "The governance, policy and institutions that enable and resource education. "
                "These institutions create the conditions for improvement at scale."
            ),
            "listItems": [
                {
                    "letter": "B",
                    "icon": "account_balance_wallet",
                    "title": "Budget allocation and resource mobilization",
                    "description": (
                        "Are institutions effectively mobilising and utilising resources "
                        "to strengthen education outcomes?"
                    ),
                },
                {
                    "letter": "R",
                    "icon": "rate_review",
                    "title": "Review, monitoring and feedback",
                    "description": (
                        "Are institutions continuously reviewing progress, learning from evidence, "
                        "and adapting their actions?"
                    ),
                },
                {
                    "letter": "I",
                    "icon": "emoji_objects",
                    "title": "Innovation and new projects",
                    "description": (
                        "Are institutions fostering innovation by initiating, experimenting, "
                        "and scaling what works?"
                    ),
                },
                {
                    "letter": "C",
                    "icon": "hub",
                    "title": "Co-creation with diverse stakeholders",
                    "description": (
                        "Are institutions meaningfully collaborating with communities, schools, "
                        "CSOs, and other stakeholders to design and implement improvement programmes?"
                    ),
                },
            ],
        },
        {
            "key": "network",
            "chipLabel": "Network",
            "diagramLabel": "Network",
            "icon": "public",
            "color": "#961c00",
            "fill": "#f3e5df",
            "diagram": {
                "shape": "full",
                "innerRadius": 216,
                "outerRadius": 304,
                "labelX": 300,
                "labelY": 58,
                "labelLayout": "inline",
                "labelAnchor": "start",
                "textOffsetX": 34,
                "textOffsetY": 4,
                "iconX": 215,
                "iconY": 45,
                "hitRadius": 344,
                "isHitTarget": True,
                "icon": {
                    "type": "material",
                    "value": "assets/icons/network.svg",
                    "color": "#961c00",
                    "width": 34,
                    "height": 34,
                },
                "text": {
                    "x": 300,
                    "y": 58,
                    "fill": "#1a1622",
                    "fontSize": 23,
                    "fontWeight": 600,
                },
            },
            "imgPath": "assets/icons/network-users.svg",
            "panelType": "story",
            "eyebrow": "Network",
            "heading": "Network",
            "subheading": "Partners and collaborators",
            "body": (
                "The wider web of actors and movements connecting everything. "
                "This layer reflects the relationships that help ideas, resources and learning move across the ecosystem."
            ),
            "frameworkNote": "Framework being defined.",
            "cta": {
                "label": "Explore Network Health",
                "link": "/network-health",
            },
            "frameworkCta": {
                "label": "Explore Network Health",
                "link": "/network-health",
            },
        },
    ],
}


def normalize_text(value):
    return re.sub(r"[^a-z0-9]", "", str(value or "").strip().lower())


def split_content(value):
    return [line.strip() for line in str(value or "").splitlines() if line.strip()]


def strip_list_marker(value):
    return re.sub(r"^[-•]\s*", "", value).strip()


def get_public_asset_base_url():
    configured_url = os.getenv("GCS_PUBLIC_BASE_URL") or os.getenv("BUCKET_URL")
    if configured_url:
        return configured_url.strip().strip('"').rstrip("/")

    bucket_name = os.getenv("BUCKET_NAME")
    if not bucket_name:
        return ""

    return f"{GCS_STORAGE_BASE_URL}/{bucket_name}"


def resolve_asset_urls(value, base_url=None):
    if base_url is None:
        base_url = get_public_asset_base_url()

    if isinstance(value, dict):
        return {key: resolve_asset_urls(item, base_url) for key, item in value.items()}

    if isinstance(value, list):
        return [resolve_asset_urls(item, base_url) for item in value]

    if isinstance(value, str) and value.startswith("assets/") and base_url:
        return f"{base_url}/{os.path.basename(value)}"

    return value


def get_layer_from_action(action):
    normalized_action = normalize_text(action)
    for action_key, layer_key in LAYER_ACTIONS.items():
        if action_key in normalized_action:
            return layer_key
    return None


def is_framework_row(section):
    return "framework" in normalize_text(section)


def build_list_items(layer_key, content):
    lines = split_content(content)
    icons = DEFAULT_LIST_ICONS.get(layer_key, [])
    list_items = []

    for index in range(0, len(lines), 2):
        title = lines[index]
        description = lines[index + 1] if index + 1 < len(lines) else ""
        list_items.append({
            "letter": title[:1].upper(),
            "icon": icons[len(list_items)] if len(list_items) < len(icons) else "article",
            "title": title,
            "description": description,
        })

    return list_items


def apply_layer_definition(layer, content):
    lines = [strip_list_marker(line) for line in split_content(content)]
    if not lines:
        return

    if layer.get("key") == "network":
        print("⚠️ Keeping default network narrative copy.")
        return

    layer["heading"] = lines[0]
    if len(lines) > 1:
        layer["subheading"] = lines[1]
    if len(lines) > 2:
        layer["body"] = " ".join(lines[2:])


def resolve_output_file(output_file):
    if os.path.isabs(output_file):
        return output_file
    return os.path.join(PROJECT_ROOT, output_file)


def get_sheet(workbook, sheet_name):
    if sheet_name in workbook.sheetnames:
        return workbook[sheet_name]

    normalized_requested = normalize_text(sheet_name)
    for name in workbook.sheetnames:
        if normalize_text(name) == normalized_requested:
            return workbook[name]

    raise KeyError(f"Sheet not found: {sheet_name}")


def generate_outcomes_model_json(excel_file, sheet_name=DEFAULT_SHEET_NAME, output_file=OUTPUT_FILE):
    output_file = resolve_output_file(output_file)
    config = copy.deepcopy(OUTCOMES_MODEL_TEMPLATE)
    layers_by_key = {layer["key"]: layer for layer in config["layers"]}

    workbook = openpyxl.load_workbook(excel_file, data_only=True)
    try:
        sheet = get_sheet(workbook, sheet_name)
    except KeyError as e:
        print(f"⚠️ {e}. Skipping outcomes model JSON generation.")
        return None

    for row in sheet.iter_rows(min_row=CONTENT_START_ROW, values_only=True):
        action = row[1] if len(row) > 1 else None
        section = row[2] if len(row) > 2 else None
        content = row[3] if len(row) > 3 else None

        layer_key = get_layer_from_action(action)
        if not layer_key or not content:
            continue

        layer = layers_by_key[layer_key]
        if is_framework_row(section):
            list_items = build_list_items(layer_key, content)
            if list_items:
                layer["listItems"] = list_items
        else:
            apply_layer_definition(layer, content)

    config = resolve_asset_urls(config)

    output_dir = os.path.dirname(output_file)
    if output_dir:
        os.makedirs(output_dir, exist_ok=True)
    with open(output_file, "w", encoding="utf-8") as json_file:
        json.dump(config, json_file, indent=2, ensure_ascii=False)

    print(f"✅ Outcomes model JSON generated: {output_file}")

    folder_url = None
    bucket_name = os.environ.get("BUCKET_NAME")

    try:
        # Dynamically import gcp_access module and upload file
        gcp_access_path = os.path.join(SCRIPT_DIR, "..", "cloud-scripts", "gcp_access.py")
        spec = importlib.util.spec_from_file_location("gcp_access", gcp_access_path)
        gcp_access = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(gcp_access)

        folder_url = gcp_access.upload_file_to_gcs_and_get_directory(
            bucket_name=bucket_name,
            source_file_path=output_file,
            destination_blob_name="sg-dashboard/outcomes-model-config.json",
        )
    except Exception as e:
        print(f"❌ Failed to upload outcomes-model-config.json to GCS. Continuing: {e}")

    if folder_url:
        print(f"Successfully uploaded and got public folder URL: {folder_url}")
    else:
        print("Failed to upload file to GCS. Check logs for details.")

    return {
        "output_file": output_file,
        "data": config,
        "gcs_path": (
            f"gs://{bucket_name}/sg-dashboard/outcomes-model-config.json"
            if bucket_name
            else ""
        ),
        "folder_url": folder_url,
    }


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Generate outcomes model JSON from Excel.")
    parser.add_argument("excel_file", help="Path to source Excel file")
    parser.add_argument("--sheet", default=DEFAULT_SHEET_NAME, help="Sheet name to read")
    parser.add_argument("--output", default=OUTPUT_FILE, help="Output JSON path")
    args = parser.parse_args()

    generate_outcomes_model_json(args.excel_file, args.sheet, args.output)

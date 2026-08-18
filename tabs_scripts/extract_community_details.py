import importlib.util
import json
import os

import openpyxl
from constants import (
    ATRISK_CHILDREN_REGULARISED,
    CHILDREN_ENROLLED,
    CHILDREN_GOT_AADHAAR,
    CHILDREN_GOT_BIRTH_CERTIFICATE,
    COMMUNITY_DETAILS_REQUIRED_COLUMNS,
    COMMUNITY_DISTRICT_COLUMN,
    COMMUNITY_DISTRICT_METRICS,
    COMMUNITY_MAP_COLUMNS,
    COMMUNITY_PIE_COLUMNS,
    COMMUNITY_STATE_COLUMN,
    DISTRICTS_ACTIVATED,
    DOCUMENTATION_COLUMNS,
    ENROLLMENT_COLUMNS,
    PAGE_METADATA,
    TOTAL_ROW_LABEL,
)

def delete_district_community_pie_charts():
    bucket_name = os.environ.get("BUCKET_NAME")

    script_dir = os.path.dirname(os.path.abspath(__file__))
    state_codes = load_state_codes()
    if not bucket_name:
        print("❌ BUCKET_NAME not set")
        return

    gcp_access_path = os.path.join(script_dir, '..', 'cloud-scripts', 'gcp_access.py')
    spec = importlib.util.spec_from_file_location('gcp_access', gcp_access_path)
    gcp_access = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(gcp_access)

    gcp_access.delete_all_district_community_pie_charts(
        bucket_name=os.environ.get("BUCKET_NAME")
    )


if __name__ == "__main__":
    delete_district_community_pie_charts()


def load_state_codes():
    script_dir = os.path.dirname(os.path.abspath(__file__))
    state_codes_file = os.path.join(script_dir, "..", "pages", "state_code_details.json")
    if not os.path.exists(state_codes_file):
        print("❌ state_code_details.json not found.")
        return None
    with open(state_codes_file, "r", encoding="utf-8") as f:
        return json.load(f)


# ✅ SAFE INT (fix for "Total")
def safe_int(val):
    try:
        return int(val)
    except:
        return 0


def load_enrollment_data(workbook):
    enrollment_data = {}
    enrollment_totals_by_state = {}
    enrollment_sum_by_state = {}

    try:
        sheet = workbook["Enrollments"]
    except KeyError:
        print("❌ 'Enrollments' sheet not found")
        return enrollment_data, enrollment_totals_by_state

    for row in sheet.iter_rows(min_row=3, values_only=True):
        state_value = row[ENROLLMENT_COLUMNS["state"]]
        district_value = row[ENROLLMENT_COLUMNS["district"]]
        state = str(state_value).strip() if state_value else None
        district = str(district_value).strip() if district_value else None

        if not state or not district:
            continue

        total_enrolled = safe_int(row[ENROLLMENT_COLUMNS["children_total"]])
        total_regularized = safe_int(row[ENROLLMENT_COLUMNS["atrisk_total"]])

        if district.lower() == TOTAL_ROW_LABEL:
            enrollment_totals_by_state[state] = {
                CHILDREN_ENROLLED: total_enrolled,
                ATRISK_CHILDREN_REGULARISED: total_regularized
            }
            continue

        enrollment_data[(state, district)] = {
            CHILDREN_ENROLLED: total_enrolled,
            ATRISK_CHILDREN_REGULARISED: total_regularized
        }

        state_sum = enrollment_sum_by_state.setdefault(state, {
            CHILDREN_ENROLLED: 0,
            ATRISK_CHILDREN_REGULARISED: 0
        })
        state_sum[CHILDREN_ENROLLED] += total_enrolled
        state_sum[ATRISK_CHILDREN_REGULARISED] += total_regularized

    for state, totals in enrollment_sum_by_state.items():
        if state not in enrollment_totals_by_state:
            enrollment_totals_by_state[state] = totals

    return enrollment_data, enrollment_totals_by_state


# ✅ UPDATED (capture TOTAL row)
def load_documentation_data(workbook):
    documentation_data = {}
    documentation_explicit_totals_by_state = {}
    documentation_sum_by_state = {}

    try:
        sheet = workbook["Documentation"]
    except KeyError:
        print("❌ 'Documentation' sheet not found")
        return documentation_data, {}, {
            CHILDREN_GOT_AADHAAR: 0,
            CHILDREN_GOT_BIRTH_CERTIFICATE: 0
        }

    for row in sheet.iter_rows(min_row=2, values_only=True):
        state_value = row[DOCUMENTATION_COLUMNS["state"]]
        district_value = row[DOCUMENTATION_COLUMNS["district"]]
        state = str(state_value).strip() if state_value else None
        district = str(district_value).strip() if district_value else None

        aadhaar = safe_int(row[DOCUMENTATION_COLUMNS["aadhaar"]])
        birth = safe_int(row[DOCUMENTATION_COLUMNS["birth_certificate"]])

        if state and district and district.lower() == TOTAL_ROW_LABEL:
            documentation_explicit_totals_by_state[state] = {
                CHILDREN_GOT_AADHAAR: aadhaar,
                CHILDREN_GOT_BIRTH_CERTIFICATE: birth
            }
            continue

        if not state or not district:
            continue

        documentation_data[(state, district)] = {
            CHILDREN_GOT_AADHAAR: aadhaar,
            CHILDREN_GOT_BIRTH_CERTIFICATE: birth
        }

        state_totals = documentation_sum_by_state.setdefault(state, {
            CHILDREN_GOT_AADHAAR: 0,
            CHILDREN_GOT_BIRTH_CERTIFICATE: 0
        })
        state_totals[CHILDREN_GOT_AADHAAR] += aadhaar
        state_totals[CHILDREN_GOT_BIRTH_CERTIFICATE] += birth

    documentation_totals_by_state = {}
    documentation_states = (
        set(documentation_sum_by_state)
        | set(documentation_explicit_totals_by_state)
    )
    for state in documentation_states:
        documentation_totals_by_state[state] = documentation_explicit_totals_by_state.get(
            state,
            documentation_sum_by_state.get(state, {
                CHILDREN_GOT_AADHAAR: 0,
                CHILDREN_GOT_BIRTH_CERTIFICATE: 0
            })
        )

    documentation_totals_global = {
        CHILDREN_GOT_AADHAAR: sum(
            totals[CHILDREN_GOT_AADHAAR]
            for totals in documentation_totals_by_state.values()
        ),
        CHILDREN_GOT_BIRTH_CERTIFICATE: sum(
            totals[CHILDREN_GOT_BIRTH_CERTIFICATE]
            for totals in documentation_totals_by_state.values()
        )
    }

    return documentation_data, documentation_totals_by_state, documentation_totals_global


def extract_community_details(excel_file):
    try:
        script_dir = os.path.dirname(os.path.abspath(__file__))
        state_codes = load_state_codes()
        if not state_codes:
            return

        workbook = openpyxl.load_workbook(excel_file, data_only=True)

        enrollment_data, enrollment_totals_by_state = load_enrollment_data(workbook)
        documentation_data, documentation_totals_by_state, documentation_totals = load_documentation_data(workbook)

        try:
            sheet = workbook[PAGE_METADATA["NEW_COMMUNITY_LED_PROGRAMS"]]
        except KeyError:
            print(f"❌ Sheet not found: {PAGE_METADATA['NEW_COMMUNITY_LED_PROGRAMS']}")
            return

        expected_headers = COMMUNITY_DETAILS_REQUIRED_COLUMNS

        column_indices = {}
        for cell in sheet[1]:
            if cell.value and str(cell.value).strip() in expected_headers:
                column_indices[str(cell.value).strip()] = cell.column

        missing_columns = [col for col in expected_headers if col not in column_indices]
        if missing_columns:
            print(f"❌ Missing required columns: {missing_columns}")
            return

        map_keys = COMMUNITY_MAP_COLUMNS
        map_display_names = {key: key for key in COMMUNITY_MAP_COLUMNS}
        pie_keys = COMMUNITY_PIE_COLUMNS

        state_data = {}

        gcp_access_path = os.path.join(script_dir, '..', 'cloud-scripts', 'gcp_access.py')
        spec = importlib.util.spec_from_file_location('gcp_access', gcp_access_path)
        gcp_access = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(gcp_access)

        for row in sheet.iter_rows(min_row=2, values_only=True):
            state_name = str(row[column_indices[COMMUNITY_STATE_COLUMN] - 1]).strip()
            district_name = str(row[column_indices[COMMUNITY_DISTRICT_COLUMN] - 1]).strip()

            if not state_name or not district_name or district_name.lower() == TOTAL_ROW_LABEL:
                continue

            if state_name not in state_codes:
                continue

            state_id = state_codes[state_name]["id"]
            district_id = state_codes[state_name].get(district_name)
            if not district_id:
                print(f"⚠️ Skipping unknown district {district_name} under state {state_name}")
                continue

            if state_id not in state_data:
                state_data[state_id] = {
                    "state_name": state_name,
                    "districts": {},
                    "overview_totals": {k: 0 for k in map_keys},
                    "enrollment_totals": {
                        CHILDREN_ENROLLED: 0,
                        ATRISK_CHILDREN_REGULARISED: 0
                    },
                    "documentation_totals": {
                        CHILDREN_GOT_AADHAAR: 0,
                        CHILDREN_GOT_BIRTH_CERTIFICATE: 0
                    },
                    "pie_totals": {k: 0 for k in pie_keys}
                }

            details = []

            for k in map_keys:
                val = safe_int(row[column_indices[k] - 1])
                state_data[state_id]["overview_totals"][k] += val
                if val != 0:
                    details.append({"value": val, "code": k})

            enroll_info = enrollment_data.get((state_name, district_name), {
                CHILDREN_ENROLLED: 0,
                ATRISK_CHILDREN_REGULARISED: 0
            })

            doc_info = documentation_data.get((state_name, district_name), {
                CHILDREN_GOT_AADHAAR: 0,
                CHILDREN_GOT_BIRTH_CERTIFICATE: 0
            })

            for label, _ in COMMUNITY_DISTRICT_METRICS:
                metric_source = enroll_info if label in enroll_info else doc_info
                if metric_source[label] != 0:
                    details.append({"value": metric_source[label], "code": label})

            state_data[state_id]["enrollment_totals"][CHILDREN_ENROLLED] += enroll_info[CHILDREN_ENROLLED]
            state_data[state_id]["enrollment_totals"][ATRISK_CHILDREN_REGULARISED] += enroll_info[ATRISK_CHILDREN_REGULARISED]

            state_data[state_id]["documentation_totals"][CHILDREN_GOT_AADHAAR] += doc_info[CHILDREN_GOT_AADHAAR]
            state_data[state_id]["documentation_totals"][CHILDREN_GOT_BIRTH_CERTIFICATE] += doc_info[CHILDREN_GOT_BIRTH_CERTIFICATE]

            state_data[state_id]["districts"][district_id] = {
                "label": district_name,
                "type": "category_1",
                "details": details
            }

            pie_totals = {}
            for k in pie_keys:
                val = row[column_indices[k] - 1] or 0
                state_data[state_id]["pie_totals"][k] += val
                pie_totals[k] = val

            district_folder = os.path.join(script_dir, "..", "districts", district_id)
            os.makedirs(district_folder, exist_ok=True)

            metrics = [
                {
                    "label": map_display_names.get(k, k),
                    "value": safe_int(row[column_indices[k] - 1]),
                    "identifier": idx
                }
                for idx, k in enumerate(map_keys, start=1)
                if safe_int(row[column_indices[k] - 1]) != 0
            ]

            for label, identifier in COMMUNITY_DISTRICT_METRICS:
                metric_source = enroll_info if label in enroll_info else doc_info
                value = metric_source[label]
                if value != 0:
                    metrics.append({"label": label, "value": value, "identifier": identifier})

            metrics_json = {"metrics": metrics}

            metrics_path = os.path.join(district_folder, "community-metrics.json")
            with open(metrics_path, "w", encoding="utf-8") as f:
                json.dump(metrics_json, f, indent=2, ensure_ascii=False)

            for fname in ["community-metrics.json"]:
                local_path = os.path.join(district_folder, fname)
                blob_path = f"sg-dashboard/districts/{district_id}/{fname}"
                folder_url = gcp_access.upload_file_to_gcs_and_get_directory(
                    bucket_name=os.environ.get("BUCKET_NAME"),
                    source_file_path=local_path,
                    destination_blob_name=blob_path
                )
                if folder_url:
                    print(f"✅ Uploaded {fname} for district {district_name} ({district_id}) to {folder_url}")
                else:
                    print(f"❌ Failed to upload {fname} for district {district_name} ({district_id})")

        for state_id, data in state_data.items():
            state_folder = os.path.join(script_dir, "..", "states", state_id)
            os.makedirs(state_folder, exist_ok=True)

            state_doc_totals = documentation_totals_by_state.get(data["state_name"], {
                CHILDREN_GOT_AADHAAR: 0,
                CHILDREN_GOT_BIRTH_CERTIFICATE: 0
            })

            state_enroll_totals = enrollment_totals_by_state.get(data["state_name"], data["enrollment_totals"])

            # build filtered overview details
            details = []

            # map_keys totals
            for k, v in data["overview_totals"].items():
                if v != 0:
                    details.append({
                        "value": v,
                        "code": map_display_names.get(k, k)
                   })

            # enrollment + documentation
            extra_fields = [
                (CHILDREN_ENROLLED, state_enroll_totals[CHILDREN_ENROLLED]),
                (ATRISK_CHILDREN_REGULARISED, state_enroll_totals[ATRISK_CHILDREN_REGULARISED]),
                (CHILDREN_GOT_AADHAAR, state_doc_totals[CHILDREN_GOT_AADHAAR]),
                (CHILDREN_GOT_BIRTH_CERTIFICATE, state_doc_totals[CHILDREN_GOT_BIRTH_CERTIFICATE]),
                (DISTRICTS_ACTIVATED, len(data["districts"]))
            ]

            for label, value in extra_fields:
                if value != 0:
                    details.append({
                        "value": value,
                        "code": label
                    })
 
            map_json = {
                "result": {
                    "districts": data["districts"],
                    "overview": {
                         "label": data["state_name"],
                         "type": "category_2",
                         "details": details
                    }
                 }
            }


            map_path = os.path.join(state_folder, "community-map.json")
            with open(map_path, "w", encoding="utf-8") as f:
                json.dump(map_json, f, indent=2, ensure_ascii=False)

            for fname in ["community-map.json"]:
                local_path = os.path.join(state_folder, fname)
                blob_path = f"sg-dashboard/states/{state_id}/{fname}"
                folder_url = gcp_access.upload_file_to_gcs_and_get_directory(
                    bucket_name=os.environ.get("BUCKET_NAME"),
                    source_file_path=local_path,
                    destination_blob_name=blob_path
                )
                if folder_url:
                    print(f"✅ Uploaded {fname} for state {data['state_name']} to {folder_url}")
                else:
                    print(f"❌ Failed to upload {fname} for state {data['state_name']}")

        state_details_path = os.path.join(script_dir, "..", "pages", "community-details-page.json")
        folder_url = gcp_access.upload_file_to_gcs_and_get_directory(
            bucket_name=os.environ.get("BUCKET_NAME"),
            source_file_path=state_details_path,
            destination_blob_name="sg-dashboard/community-details-page.json"
        )

        if folder_url:
            print(f"✅ Uploaded community-details-page.json to {folder_url}")
        else:
            print("❌ Failed to upload community-details-page.json")

    except Exception as e:
        print(f"❌ Error: {str(e)}")


if __name__ == "__main__":
    import sys
    if len(sys.argv) < 2:
        print("Usage: python extract_community_details.py <excel_file>")
    else:
        extract_community_details(sys.argv[1])
